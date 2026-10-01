"""Servidor de ferramentas do assistente (OpenAPI).

Usado de dois jeitos:
  * pelo Open WebUI, como "tool server" (cada rota com operation_id vira uma ferramenta);
  * pelo modo de voz avançado, que chama as mesmas rotas.

Também expõe, fora do esquema de ferramentas:
  * POST /v1/images/generations — API de imagens no formato do OpenAI, que o
    Open WebUI usa como motor de imagens (por trás, ComfyUI);
  * POST /extrair_texto — lê PDF/DOCX/TXT enviados no modo de voz;
  * GET  /arquivos/<nome> — downloads dos documentos e imagens gerados.
"""
from __future__ import annotations

import asyncio
import base64
import datetime as dt
import io
import json
import logging
import os
import random
import re
import time
import unicodedata
import uuid
from pathlib import Path
from typing import Literal
from zoneinfo import ZoneInfo

import httpx
from fastapi import Depends, FastAPI, File, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

import documentos

log = logging.getLogger("ferramentas")
logging.basicConfig(level=os.getenv("NIVEL_LOG", "INFO"), format="%(asctime)s %(levelname)s %(name)s: %(message)s")

TOKEN = os.getenv("FERRAMENTAS_TOKEN", "")
SEARXNG_URL = os.getenv("SEARXNG_URL", "http://searxng:8080")
COMFYUI_URL = os.getenv("COMFYUI_URL", "http://comfyui:8188")
OLLAMA_URL = os.getenv("OLLAMA_URL", "http://ollama:11434")
LLM_BACKEND = os.getenv("LLM_BACKEND", "ollama")
SD_MODELO = os.getenv("SD_MODELO", "sdxl")
IMAGEM_PASSOS = int(os.getenv("IMAGEM_PASSOS", "25"))
LIBERAR_VRAM = os.getenv("LIBERAR_VRAM_PARA_IMAGEM", "0") == "1"
PASTA = Path(os.getenv("ARQUIVOS_DIR", "/dados/arquivos"))
URL_PUBLICA = os.getenv("ARQUIVOS_URL_PUBLICA", "http://localhost:3002").rstrip("/")
FUSO = os.getenv("FUSO_HORARIO", "America/Sao_Paulo")
NAVEGADOR = "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/140 Safari/537.36"

CHECKPOINTS = {
    "sdxl": {"arquivo": "sd_xl_base_1.0.safetensors", "lado": 1024, "cfg": 6.5, "amostrador": "dpmpp_2m", "agenda": "karras"},
    "sd15": {"arquivo": "v1-5-pruned-emaonly-fp16.safetensors", "lado": 512, "cfg": 7.0, "amostrador": "dpmpp_2m", "agenda": "karras"},
}
NEGATIVO_PADRAO = "lowres, blurry, bad anatomy, bad hands, extra fingers, deformed, watermark, text, signature, jpeg artifacts"

PASTA.mkdir(parents=True, exist_ok=True)
seguranca = HTTPBearer(auto_error=False)


def autenticar(cred: HTTPAuthorizationCredentials | None = Depends(seguranca)) -> None:
    if TOKEN and (cred is None or cred.credentials != TOKEN):
        raise HTTPException(401, "token inválido")


app = FastAPI(
    title="Ferramentas do assistente",
    version="1.0",
    description="Busca na web, leitura de páginas, geração de imagens, criação de documentos e data/hora.",
)
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])
app.mount("/arquivos", StaticFiles(directory=str(PASTA)), name="arquivos")


def _slug(texto: str, limite: int = 50) -> str:
    t = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    t = re.sub(r"[^a-zA-Z0-9]+", "-", t).strip("-").lower()
    return (t[:limite].strip("-") or "arquivo")


def _url(nome: str) -> str:
    return f"{URL_PUBLICA}/arquivos/{nome}"


# ----------------------------------------------------------------------------
# Busca na web
# ----------------------------------------------------------------------------
class PedidoBusca(BaseModel):
    consulta: str = Field(..., description="O que pesquisar, em linguagem natural (de preferência em português).")
    quantidade: int = Field(5, ge=1, le=10, description="Quantos resultados trazer (1 a 10).")


class Resultado(BaseModel):
    titulo: str
    url: str
    trecho: str


class RespostaBusca(BaseModel):
    consulta: str
    resultados: list[Resultado]
    resposta_direta: str | None = None


@app.post("/buscar_web", operation_id="buscar_web", response_model=RespostaBusca, dependencies=[Depends(autenticar)],
          summary="Pesquisa na internet",
          description="Pesquisa na internet (via SearXNG local) e devolve títulos, links e trechos. "
                      "Use para notícias, preços, clima, fatos recentes ou qualquer coisa que você não saiba com certeza.")
async def buscar_web(p: PedidoBusca) -> RespostaBusca:
    params = {"q": p.consulta, "format": "json", "language": "pt-BR", "safesearch": "1"}
    async with httpx.AsyncClient(timeout=20) as c:
        try:
            r = await c.get(f"{SEARXNG_URL}/search", params=params)
            r.raise_for_status()
            dados = r.json()
        except Exception as e:  # noqa: BLE001
            log.warning("SearXNG falhou: %s", e)
            raise HTTPException(502, f"a busca falhou: {e}") from e
    resultados = []
    vistos = set()
    for item in dados.get("results", []):
        url = item.get("url", "")
        if not url or url in vistos:
            continue
        vistos.add(url)
        resultados.append(Resultado(titulo=item.get("title", "")[:200], url=url, trecho=(item.get("content") or "")[:500]))
        if len(resultados) >= p.quantidade:
            break
    direta = None
    if dados.get("answers"):
        a = dados["answers"][0]
        direta = a.get("answer") if isinstance(a, dict) else str(a)
    elif dados.get("infoboxes"):
        direta = (dados["infoboxes"][0].get("content") or "")[:600] or None
    return RespostaBusca(consulta=p.consulta, resultados=resultados, resposta_direta=direta)


class PedidoPagina(BaseModel):
    url: str = Field(..., description="Endereço completo da página (http ou https).")


class RespostaPagina(BaseModel):
    url: str
    titulo: str
    texto: str


def _texto_de_html(conteudo: str) -> tuple[str, str]:
    from bs4 import BeautifulSoup

    sopa = BeautifulSoup(conteudo, "html.parser")
    titulo = (sopa.title.string or "").strip() if sopa.title and sopa.title.string else ""
    for tag in sopa(["script", "style", "noscript", "nav", "footer", "header", "aside", "form", "svg", "iframe"]):
        tag.decompose()
    principal = sopa.find("article") or sopa.find("main") or sopa.body or sopa
    texto = principal.get_text("\n", strip=True)
    texto = re.sub(r"\n{2,}", "\n", texto)
    return titulo, texto


def _texto_de_pdf(dados: bytes) -> str:
    from pypdf import PdfReader

    leitor = PdfReader(io.BytesIO(dados))
    return "\n".join((p.extract_text() or "") for p in leitor.pages[:60])


@app.post("/ler_pagina", operation_id="ler_pagina", response_model=RespostaPagina, dependencies=[Depends(autenticar)],
          summary="Lê o conteúdo de uma página da internet",
          description="Abre um link (página ou PDF) e devolve o texto principal. Use depois de buscar_web quando os trechos não bastarem.")
async def ler_pagina(p: PedidoPagina) -> RespostaPagina:
    if not re.match(r"^https?://", p.url):
        raise HTTPException(400, "URL deve começar com http:// ou https://")
    async with httpx.AsyncClient(timeout=25, follow_redirects=True, headers={"User-Agent": NAVEGADOR, "Accept-Language": "pt-BR,pt;q=0.9"}) as c:
        try:
            r = await c.get(p.url)
            r.raise_for_status()
        except Exception as e:  # noqa: BLE001
            raise HTTPException(502, f"não consegui abrir a página: {e}") from e
    tipo = r.headers.get("content-type", "")
    if "pdf" in tipo or p.url.lower().endswith(".pdf"):
        titulo, texto = Path(p.url).name, await asyncio.to_thread(_texto_de_pdf, r.content)
    else:
        titulo, texto = await asyncio.to_thread(_texto_de_html, r.text)
    return RespostaPagina(url=str(r.url), titulo=titulo, texto=texto[:8000])


# ----------------------------------------------------------------------------
# Imagens (ComfyUI)
# ----------------------------------------------------------------------------
def _fluxo_comfyui(prompt: str, negativo: str, largura: int, altura: int, passos: int, semente: int) -> dict:
    ck = CHECKPOINTS.get(SD_MODELO, CHECKPOINTS["sd15"])
    return {
        "4": {"class_type": "CheckpointLoaderSimple", "inputs": {"ckpt_name": ck["arquivo"]}},
        "6": {"class_type": "CLIPTextEncode", "inputs": {"text": prompt, "clip": ["4", 1]}},
        "7": {"class_type": "CLIPTextEncode", "inputs": {"text": negativo, "clip": ["4", 1]}},
        "5": {"class_type": "EmptyLatentImage", "inputs": {"width": largura, "height": altura, "batch_size": 1}},
        "3": {
            "class_type": "KSampler",
            "inputs": {
                "seed": semente, "steps": passos, "cfg": ck["cfg"], "sampler_name": ck["amostrador"],
                "scheduler": ck["agenda"], "denoise": 1.0,
                "model": ["4", 0], "positive": ["6", 0], "negative": ["7", 0], "latent_image": ["5", 0],
            },
        },
        "8": {"class_type": "VAEDecode", "inputs": {"samples": ["3", 0], "vae": ["4", 2]}},
        "9": {"class_type": "SaveImage", "inputs": {"filename_prefix": "assistente", "images": ["8", 0]}},
    }


def _dimensoes(formato: str | None = None, tamanho: str | None = None) -> tuple[int, int]:
    lado = CHECKPOINTS.get(SD_MODELO, CHECKPOINTS["sd15"])["lado"]
    if tamanho and re.fullmatch(r"\d+x\d+", tamanho):
        w, h = (int(x) for x in tamanho.split("x"))
        # Mantém a proporção pedida, mas na resolução nativa do modelo
        escala = (lado * lado / max(w * h, 1)) ** 0.5
        w, h = int(w * escala), int(h * escala)
    elif formato == "retrato":
        w, h = int(lado * 0.8125), int(lado * 1.1875)
    elif formato == "paisagem":
        w, h = int(lado * 1.1875), int(lado * 0.8125)
    else:
        w = h = lado
    return max(256, w // 64 * 64), max(256, h // 64 * 64)


async def _liberar_llm(c: httpx.AsyncClient) -> None:
    """Descarrega o LLM do Ollama para o Stable Diffusion caber na VRAM. Ele recarrega sozinho depois."""
    if LLM_BACKEND != "ollama":
        return
    try:
        r = await c.get(f"{OLLAMA_URL}/api/ps", timeout=5)
        for m in r.json().get("models", []):
            await c.post(f"{OLLAMA_URL}/api/generate", json={"model": m["name"], "keep_alive": 0}, timeout=30)
            log.info("LLM %s descarregado para gerar imagem", m["name"])
    except Exception as e:  # noqa: BLE001
        log.warning("Não consegui descarregar o LLM: %s", e)


async def gerar_png(prompt: str, largura: int, altura: int, negativo: str = NEGATIVO_PADRAO, passos: int | None = None) -> bytes:
    if SD_MODELO == "nenhum":
        raise HTTPException(503, "geração de imagens desativada nesta máquina")
    async with httpx.AsyncClient(timeout=30) as c:
        if LIBERAR_VRAM:
            await _liberar_llm(c)
        fluxo = _fluxo_comfyui(prompt, negativo, largura, altura, passos or IMAGEM_PASSOS, random.randint(1, 2**31 - 1))
        try:
            r = await c.post(f"{COMFYUI_URL}/prompt", json={"prompt": fluxo, "client_id": uuid.uuid4().hex})
            r.raise_for_status()
        except httpx.HTTPStatusError as e:
            raise HTTPException(502, f"ComfyUI recusou o pedido: {e.response.text[:400]}") from e
        except Exception as e:  # noqa: BLE001
            raise HTTPException(502, f"ComfyUI indisponível: {e}") from e
        prompt_id = r.json()["prompt_id"]
        limite = time.time() + float(os.getenv("IMAGEM_TEMPO_MAXIMO", "900"))
        saida = None
        while time.time() < limite:
            await asyncio.sleep(0.5)
            h = (await c.get(f"{COMFYUI_URL}/history/{prompt_id}")).json()
            item = h.get(prompt_id)
            if not item:
                continue
            status = item.get("status", {})
            if status.get("status_str") == "error":
                msgs = [m for m in status.get("messages", []) if m and m[0] == "execution_error"]
                detalhe = msgs[0][1].get("exception_message", "") if msgs else ""
                raise HTTPException(500, f"ComfyUI falhou ao gerar a imagem: {detalhe[:300]}")
            imagens = item.get("outputs", {}).get("9", {}).get("images", [])
            if imagens:
                saida = imagens[0]
                break
        if saida is None:
            raise HTTPException(504, "a geração da imagem demorou demais")
        img = await c.get(f"{COMFYUI_URL}/view", params={"filename": saida["filename"], "subfolder": saida.get("subfolder", ""), "type": saida.get("type", "output")})
        img.raise_for_status()
        if LIBERAR_VRAM:
            # Devolve a VRAM ao LLM
            try:
                await c.post(f"{COMFYUI_URL}/free", json={"unload_models": True, "free_memory": True})
            except Exception:  # noqa: BLE001
                pass
        return img.content


class PedidoImagem(BaseModel):
    descricao_em_ingles: str = Field(
        ..., description="Descrição detalhada da imagem EM INGLÊS (o Stable Diffusion entende inglês): assunto, estilo, "
                         "iluminação, enquadramento. Ex.: 'a cute orange cat astronaut floating in space, digital art, detailed'."
    )
    formato: Literal["quadrado", "retrato", "paisagem"] = Field("quadrado", description="Proporção da imagem.")


class RespostaImagem(BaseModel):
    url: str
    markdown: str
    segundos: float


@app.post("/gerar_imagem", operation_id="gerar_imagem", response_model=RespostaImagem, dependencies=[Depends(autenticar)],
          summary="Gera uma imagem com Stable Diffusion",
          description="Cria uma imagem a partir de uma descrição em inglês e devolve o link. Mostre o markdown devolvido para o usuário ver a imagem.")
async def gerar_imagem(p: PedidoImagem) -> RespostaImagem:
    t0 = time.time()
    w, h = _dimensoes(p.formato)
    png = await gerar_png(p.descricao_em_ingles, w, h)
    nome = f"imagem-{_slug(p.descricao_em_ingles, 40)}-{uuid.uuid4().hex[:8]}.png"
    (PASTA / nome).write_bytes(png)
    return RespostaImagem(url=_url(nome), markdown=f"![{p.descricao_em_ingles[:80]}]({_url(nome)})", segundos=round(time.time() - t0, 1))


@app.post("/v1/images/generations", include_in_schema=False, dependencies=[Depends(autenticar)])
async def imagens_openai(request: Request) -> dict:
    corpo = await request.json()
    prompt = str(corpo.get("prompt", "")).strip()
    if not prompt:
        raise HTTPException(400, "prompt vazio")
    n = max(1, min(int(corpo.get("n") or 1), 4))
    w, h = _dimensoes(tamanho=corpo.get("size"))
    passos = corpo.get("steps")
    dados = []
    for _ in range(n):
        png = await gerar_png(prompt, w, h, passos=int(passos) if passos else None)
        dados.append({"b64_json": base64.b64encode(png).decode(), "revised_prompt": prompt})
    return {"created": int(time.time()), "data": dados}


@app.get("/v1/models", include_in_schema=False)
async def modelos_imagem() -> dict:
    return {"object": "list", "data": [{"id": SD_MODELO, "object": "model"}]}


# ----------------------------------------------------------------------------
# Documentos
# ----------------------------------------------------------------------------
class PedidoDocumento(BaseModel):
    titulo: str = Field(..., description="Título do documento.")
    conteudo_markdown: str = Field(..., description="Conteúdo completo em markdown (títulos com #, listas com -, tabelas com |).")
    formato: Literal["pdf", "txt", "md", "docx"] = Field("pdf", description="Formato do arquivo: pdf, txt, md ou docx (Word).")


class RespostaDocumento(BaseModel):
    url: str
    nome_arquivo: str
    formato: str
    bytes: int
    markdown: str


@app.post("/criar_documento", operation_id="criar_documento", response_model=RespostaDocumento, dependencies=[Depends(autenticar)],
          summary="Cria um documento para download (PDF, TXT, Markdown ou Word)",
          description="Gera um arquivo a partir de um conteúdo em markdown e devolve o link de download. Sempre mostre o link ao usuário.")
async def criar_documento(p: PedidoDocumento) -> RespostaDocumento:
    nome = f"{_slug(p.titulo)}-{uuid.uuid4().hex[:6]}.{p.formato}"
    destino = PASTA / nome
    try:
        if p.formato == "pdf":
            await asyncio.to_thread(documentos.para_pdf, p.titulo, p.conteudo_markdown, destino)
        elif p.formato == "docx":
            await asyncio.to_thread(documentos.para_docx, p.titulo, p.conteudo_markdown, destino)
        elif p.formato == "txt":
            destino.write_text(documentos.para_txt(p.titulo, p.conteudo_markdown), encoding="utf-8")
        else:
            destino.write_text(documentos.para_md(p.titulo, p.conteudo_markdown), encoding="utf-8")
    except Exception as e:  # noqa: BLE001
        log.exception("Falha ao criar documento")
        raise HTTPException(500, f"falha ao criar o documento: {e}") from e
    return RespostaDocumento(
        url=_url(nome), nome_arquivo=nome, formato=p.formato, bytes=destino.stat().st_size,
        markdown=f"[📄 Baixar {nome}]({_url(nome)})",
    )


# ----------------------------------------------------------------------------
# Data e hora
# ----------------------------------------------------------------------------
DIAS = ["segunda-feira", "terça-feira", "quarta-feira", "quinta-feira", "sexta-feira", "sábado", "domingo"]
MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]


class RespostaDataHora(BaseModel):
    data_hora: str
    iso: str
    fuso: str


@app.get("/data_hora", operation_id="data_hora", response_model=RespostaDataHora, dependencies=[Depends(autenticar)],
         summary="Data e hora atuais", description="Devolve a data e a hora atuais no fuso do usuário.")
async def data_hora() -> RespostaDataHora:
    agora = dt.datetime.now(ZoneInfo(FUSO))
    texto = f"{DIAS[agora.weekday()]}, {agora.day} de {MESES[agora.month - 1]} de {agora.year}, {agora:%H:%M}"
    return RespostaDataHora(data_hora=texto, iso=agora.isoformat(timespec="seconds"), fuso=FUSO)


# ----------------------------------------------------------------------------
# Leitura de arquivos enviados (modo de voz)
# ----------------------------------------------------------------------------
@app.post("/extrair_texto", include_in_schema=False, dependencies=[Depends(autenticar)])
async def extrair_texto(arquivo: UploadFile = File(...)) -> dict:
    dados = await arquivo.read()
    nome = arquivo.filename or "arquivo"
    ext = Path(nome).suffix.lower()
    try:
        if ext == ".pdf":
            texto = await asyncio.to_thread(_texto_de_pdf, dados)
        elif ext == ".docx":
            from docx import Document

            d = Document(io.BytesIO(dados))
            partes = [p.text for p in d.paragraphs]
            for t in d.tables:
                for linha in t.rows:
                    partes.append(" | ".join(c.text for c in linha.cells))
            texto = "\n".join(partes)
        elif ext in (".html", ".htm"):
            texto = _texto_de_html(dados.decode("utf-8", "ignore"))[1]
        elif ext == ".json":
            texto = json.dumps(json.loads(dados), ensure_ascii=False, indent=1)
        else:
            texto = dados.decode("utf-8", "ignore")
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, f"não consegui ler {nome}: {e}") from e
    texto = texto.strip()
    return {"nome": nome, "caracteres": len(texto), "texto": texto[:30000], "truncado": len(texto) > 30000}


@app.get("/health", include_in_schema=False)
async def saude() -> dict:
    estado = {"ferramentas": "ok", "sd_modelo": SD_MODELO, "liberar_vram_para_imagem": LIBERAR_VRAM}
    async with httpx.AsyncClient(timeout=3) as c:
        for nome, url in (("searxng", f"{SEARXNG_URL}/healthz"), ("comfyui", f"{COMFYUI_URL}/system_stats")):
            try:
                estado[nome] = "ok" if (await c.get(url)).status_code == 200 else "erro"
            except Exception:  # noqa: BLE001
                estado[nome] = "indisponível"
    return estado
