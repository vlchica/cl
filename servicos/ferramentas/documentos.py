"""Converte markdown simples em PDF, DOCX, TXT ou MD.

Cobre o que um assistente costuma escrever: títulos, parágrafos, listas,
listas numeradas, citações, blocos de código, tabelas e negrito/itálico/código/links.
"""
from __future__ import annotations

import html
import re
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Bloco:
    tipo: str  # titulo, paragrafo, lista, numerada, citacao, codigo, tabela, linha
    texto: str = ""
    nivel: int = 0
    itens: list[str] = field(default_factory=list)
    linhas: list[list[str]] = field(default_factory=list)


def _celulas(linha: str) -> list[str]:
    return [c.strip() for c in linha.strip().strip("|").split("|")]


def analisar(md: str) -> list[Bloco]:
    linhas = (md or "").replace("\r\n", "\n").split("\n")
    blocos: list[Bloco] = []
    i = 0
    paragrafo: list[str] = []

    def fechar_paragrafo() -> None:
        if paragrafo:
            blocos.append(Bloco("paragrafo", " ".join(s.strip() for s in paragrafo)))
            paragrafo.clear()

    while i < len(linhas):
        linha = linhas[i]
        s = linha.strip()
        if s.startswith("```"):
            fechar_paragrafo()
            codigo = []
            i += 1
            while i < len(linhas) and not linhas[i].strip().startswith("```"):
                codigo.append(linhas[i])
                i += 1
            blocos.append(Bloco("codigo", "\n".join(codigo)))
            i += 1
            continue
        if not s:
            fechar_paragrafo()
            i += 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            fechar_paragrafo()
            blocos.append(Bloco("titulo", m.group(2).strip().rstrip("#").strip(), nivel=len(m.group(1))))
            i += 1
            continue
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", s):
            fechar_paragrafo()
            blocos.append(Bloco("linha"))
            i += 1
            continue
        if s.startswith("|") and i + 1 < len(linhas) and re.fullmatch(r"\|?\s*:?-{2,}.*", linhas[i + 1].strip()):
            fechar_paragrafo()
            tabela = [_celulas(s)]
            i += 2
            while i < len(linhas) and linhas[i].strip().startswith("|"):
                tabela.append(_celulas(linhas[i]))
                i += 1
            blocos.append(Bloco("tabela", linhas=tabela))
            continue
        if re.match(r"^[-*+•]\s+", s):
            fechar_paragrafo()
            itens = []
            while i < len(linhas) and re.match(r"^\s*[-*+•]\s+", linhas[i]):
                itens.append(re.sub(r"^\s*[-*+•]\s+", "", linhas[i]).strip())
                i += 1
            blocos.append(Bloco("lista", itens=itens))
            continue
        if re.match(r"^\d+[.)]\s+", s):
            fechar_paragrafo()
            itens = []
            while i < len(linhas) and re.match(r"^\s*\d+[.)]\s+", linhas[i]):
                itens.append(re.sub(r"^\s*\d+[.)]\s+", "", linhas[i]).strip())
                i += 1
            blocos.append(Bloco("numerada", itens=itens))
            continue
        if s.startswith(">"):
            fechar_paragrafo()
            citacao = []
            while i < len(linhas) and linhas[i].strip().startswith(">"):
                citacao.append(linhas[i].strip().lstrip(">").strip())
                i += 1
            blocos.append(Bloco("citacao", " ".join(citacao)))
            continue
        paragrafo.append(linha)
        i += 1
    fechar_paragrafo()
    return blocos


def sem_markdown(t: str) -> str:
    t = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", t)
    t = re.sub(r"\[([^\]]+)\]\(([^)]*)\)", r"\1 (\2)", t)
    t = re.sub(r"(\*\*|__)(.+?)\1", r"\2", t)
    t = re.sub(r"(?<!\w)(\*|_)(.+?)\1(?!\w)", r"\2", t)
    return t.replace("`", "")


# ----------------------------------------------------------------------------
# TXT e MD
# ----------------------------------------------------------------------------
def para_txt(titulo: str, md: str) -> str:
    blocos = analisar(md)
    if blocos and blocos[0].tipo == "titulo" and blocos[0].nivel == 1:
        titulo = sem_markdown(blocos.pop(0).texto)
    saida = [titulo, "=" * min(len(titulo), 80), ""]
    for b in blocos:
        if b.tipo == "titulo":
            t = sem_markdown(b.texto)
            saida += ["", t, ("=" if b.nivel <= 2 else "-") * min(len(t), 80)]
        elif b.tipo == "paragrafo":
            saida += [sem_markdown(b.texto), ""]
        elif b.tipo == "lista":
            saida += [f"  • {sem_markdown(x)}" for x in b.itens] + [""]
        elif b.tipo == "numerada":
            saida += [f"  {n}. {sem_markdown(x)}" for n, x in enumerate(b.itens, 1)] + [""]
        elif b.tipo == "citacao":
            saida += [f"  “{sem_markdown(b.texto)}”", ""]
        elif b.tipo == "codigo":
            saida += ["    " + l for l in b.texto.split("\n")] + [""]
        elif b.tipo == "tabela":
            colunas = max(len(l) for l in b.linhas)
            linhas = [[sem_markdown(l[c]) if c < len(l) else "" for c in range(colunas)] for l in b.linhas]
            larguras = [max(len(l[c]) for l in linhas) for c in range(colunas)]
            for n, l in enumerate(linhas):
                saida.append(" | ".join(c.ljust(w) for c, w in zip(l, larguras)).rstrip())
                if n == 0:
                    saida.append("-+-".join("-" * w for w in larguras))
            saida.append("")
        elif b.tipo == "linha":
            saida += ["-" * 40, ""]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(saida)).strip() + "\n"


def para_md(titulo: str, md: str) -> str:
    corpo = md.strip()
    if not corpo.lstrip().startswith("# "):
        corpo = f"# {titulo}\n\n{corpo}"
    return corpo + "\n"


# ----------------------------------------------------------------------------
# PDF (reportlab)
# ----------------------------------------------------------------------------
FONTES_DEJAVU = Path("/usr/share/fonts/truetype/dejavu")


def _fontes_pdf() -> tuple[str, str, str, str]:
    """Usa DejaVu (cobre acentos, símbolos e travessões); sem ela, Helvetica."""
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if (FONTES_DEJAVU / "DejaVuSans.ttf").exists():
        if "DejaVuSans" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("DejaVuSans", str(FONTES_DEJAVU / "DejaVuSans.ttf")))
            pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", str(FONTES_DEJAVU / "DejaVuSans-Bold.ttf")))
            obliqua = FONTES_DEJAVU / "DejaVuSans-Oblique.ttf"
            negrito_obliqua = FONTES_DEJAVU / "DejaVuSans-BoldOblique.ttf"
            pdfmetrics.registerFont(TTFont("DejaVuSans-Oblique", str(obliqua if obliqua.exists() else FONTES_DEJAVU / "DejaVuSans.ttf")))
            pdfmetrics.registerFont(TTFont("DejaVuSans-BoldOblique", str(negrito_obliqua if negrito_obliqua.exists() else FONTES_DEJAVU / "DejaVuSans-Bold.ttf")))
            pdfmetrics.registerFont(TTFont("DejaVuSansMono", str(FONTES_DEJAVU / "DejaVuSansMono.ttf")))
            addMapping("DejaVuSans", 0, 0, "DejaVuSans")
            addMapping("DejaVuSans", 1, 0, "DejaVuSans-Bold")
            addMapping("DejaVuSans", 0, 1, "DejaVuSans-Oblique")
            addMapping("DejaVuSans", 1, 1, "DejaVuSans-BoldOblique")
        return "DejaVuSans", "DejaVuSans-Bold", "DejaVuSansMono", "DejaVuSans-Oblique"
    return "Helvetica", "Helvetica-Bold", "Courier", "Helvetica-Oblique"


def _inline_pdf(t: str, mono: str) -> str:
    t = html.escape(t, quote=False)
    t = re.sub(r"`([^`]+)`", rf'<font face="{mono}">\1</font>', t)
    t = re.sub(r"\[([^\]]+)\]\(([^)\s]+)\)", r'<link href="\2" color="#1a5fb4"><u>\1</u></link>', t)
    t = re.sub(r"(\*\*|__)(.+?)\1", r"<b>\2</b>", t)
    t = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", r"<i>\1</i>", t)
    t = re.sub(r"(?<!\w)_(?!\s)(.+?)(?<!\s)_(?!\w)", r"<i>\1</i>", t)
    return t


def para_pdf(titulo: str, md: str, destino: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        HRFlowable, ListFlowable, ListItem, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    normal, negrito, mono, _italico = _fontes_pdf()
    base = getSampleStyleSheet()
    estilos = {
        "corpo": ParagraphStyle("corpo", parent=base["BodyText"], fontName=normal, fontSize=11, leading=15.5, alignment=TA_LEFT, spaceAfter=6),
        "titulo_doc": ParagraphStyle("titulo_doc", parent=base["Title"], fontName=negrito, fontSize=20, leading=24, spaceAfter=14),
        "citacao": ParagraphStyle("citacao", parent=base["BodyText"], fontName=normal, fontSize=11, leading=15, leftIndent=18, textColor=colors.HexColor("#444444"), borderPadding=4),
        "codigo": ParagraphStyle("codigo", parent=base["Code"], fontName=mono, fontSize=9, leading=12, backColor=colors.HexColor("#f4f4f4"), borderPadding=6, leftIndent=6),
        "celula": ParagraphStyle("celula", parent=base["BodyText"], fontName=normal, fontSize=9.5, leading=12),
    }
    tamanhos = {1: 17, 2: 15, 3: 13, 4: 12, 5: 11, 6: 11}

    doc = SimpleDocTemplate(
        str(destino), pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
        title=titulo, author="Assistente local", lang="pt-BR",
    )
    hist = []
    blocos = analisar(md)
    if not (blocos and blocos[0].tipo == "titulo" and blocos[0].nivel == 1):
        hist.append(Paragraph(_inline_pdf(titulo, mono), estilos["titulo_doc"]))
    for b in blocos:
        if b.tipo == "titulo":
            if b.nivel == 1:
                hist.append(Paragraph(_inline_pdf(b.texto, mono), estilos["titulo_doc"]))
            else:
                est = ParagraphStyle(f"h{b.nivel}", parent=base["Heading2"], fontName=negrito, fontSize=tamanhos[b.nivel], leading=tamanhos[b.nivel] + 4, spaceBefore=10, spaceAfter=6)
                hist.append(Paragraph(_inline_pdf(b.texto, mono), est))
        elif b.tipo == "paragrafo":
            hist.append(Paragraph(_inline_pdf(b.texto, mono), estilos["corpo"]))
        elif b.tipo in ("lista", "numerada"):
            itens = [ListItem(Paragraph(_inline_pdf(x, mono), estilos["corpo"]), leftIndent=14) for x in b.itens]
            if b.tipo == "lista":
                hist.append(ListFlowable(itens, bulletType="bullet", start="•", leftIndent=14, bulletFontName=normal))
            else:
                hist.append(ListFlowable(itens, bulletType="1", bulletFormat="%s.", leftIndent=18, bulletFontName=normal))
            hist.append(Spacer(1, 4))
        elif b.tipo == "citacao":
            hist.append(Paragraph(_inline_pdf(b.texto, mono), estilos["citacao"]))
        elif b.tipo == "codigo":
            hist.append(Preformatted(b.texto, estilos["codigo"]))
            hist.append(Spacer(1, 6))
        elif b.tipo == "tabela":
            colunas = max(len(l) for l in b.linhas)
            dados = [[Paragraph(_inline_pdf(l[c] if c < len(l) else "", mono), estilos["celula"]) for c in range(colunas)] for l in b.linhas]
            t = Table(dados, repeatRows=1, hAlign="LEFT")
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef7")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#b0b8c4")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTNAME", (0, 0), (-1, 0), negrito),
            ]))
            hist += [t, Spacer(1, 8)]
        elif b.tipo == "linha":
            hist.append(HRFlowable(width="100%", color=colors.HexColor("#cccccc"), spaceBefore=6, spaceAfter=6))

    def rodape(canvas, documento):
        canvas.saveState()
        canvas.setFont(normal, 8)
        canvas.setFillColor(colors.HexColor("#888888"))
        canvas.drawRightString(A4[0] - 2.2 * cm, 1.2 * cm, f"{titulo} — página {documento.page}")
        canvas.restoreState()

    doc.build(hist, onFirstPage=rodape, onLaterPages=rodape)


# ----------------------------------------------------------------------------
# DOCX (python-docx)
# ----------------------------------------------------------------------------
def _runs_docx(paragrafo, texto: str) -> None:
    padrao = re.compile(r"(\*\*.+?\*\*|__.+?__|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<!\w)\*[^*\s][^*]*?\*(?!\w)|(?<!\w)_[^_\s][^_]*?_(?!\w))")
    pos = 0
    for m in padrao.finditer(texto):
        if m.start() > pos:
            paragrafo.add_run(texto[pos : m.start()])
        trecho = m.group(0)
        if trecho.startswith(("**", "__")):
            paragrafo.add_run(trecho[2:-2]).bold = True
        elif trecho.startswith("`"):
            r = paragrafo.add_run(trecho[1:-1])
            r.font.name = "Consolas"
        elif trecho.startswith("["):
            mm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", trecho)
            paragrafo.add_run(f"{mm.group(1)} ({mm.group(2)})")
        else:
            paragrafo.add_run(trecho[1:-1]).italic = True
        pos = m.end()
    if pos < len(texto):
        paragrafo.add_run(texto[pos:])


def para_docx(titulo: str, md: str, destino: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    d = Document()
    d.core_properties.title = titulo
    d.core_properties.language = "pt-BR"
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    blocos = analisar(md)
    if not (blocos and blocos[0].tipo == "titulo" and blocos[0].nivel == 1):
        d.add_heading(titulo, level=0)
    for b in blocos:
        if b.tipo == "titulo":
            d.add_heading(sem_markdown(b.texto), level=0 if b.nivel == 1 else min(b.nivel - 1, 9))
        elif b.tipo == "paragrafo":
            _runs_docx(d.add_paragraph(), b.texto)
        elif b.tipo == "lista":
            for x in b.itens:
                _runs_docx(d.add_paragraph(style="List Bullet"), x)
        elif b.tipo == "numerada":
            for x in b.itens:
                _runs_docx(d.add_paragraph(style="List Number"), x)
        elif b.tipo == "citacao":
            _runs_docx(d.add_paragraph(style="Quote"), b.texto)
        elif b.tipo == "codigo":
            p = d.add_paragraph()
            r = p.add_run(b.texto)
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        elif b.tipo == "tabela":
            colunas = max(len(l) for l in b.linhas)
            t = d.add_table(rows=len(b.linhas), cols=colunas)
            t.style = "Table Grid"
            for i, linha in enumerate(b.linhas):
                for j in range(colunas):
                    celula = t.cell(i, j)
                    celula.text = ""
                    _runs_docx(celula.paragraphs[0], linha[j] if j < len(linha) else "")
                    if i == 0:
                        for r in celula.paragraphs[0].runs:
                            r.bold = True
        elif b.tipo == "linha":
            d.add_paragraph("―" * 30)
    d.save(str(destino))
